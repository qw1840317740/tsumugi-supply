"""Import the 2026-10-06 quote workbooks into the public product catalog.

The importer is intentionally conservative: discontinued rows are excluded,
JAN is the deduplication key, and embedded workbook images are exported only
for the matching product row. Running it again is safe.
"""

from __future__ import annotations

import io
import json
import re
import unicodedata
from pathlib import Path

import openpyxl
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "assets" / "js" / "data.js"
PRODUCT_IMAGE_DIR = ROOT / "assets" / "products"
QUOTE_FILE = Path(r"C:\Users\download\Desktop\1006 - お見積り依頼（開成）26.10.6.xlsx")
AKAIBOHSHI_FILE = Path(r"C:\Users\download\Desktop\赤い帽子 开成26.10.6.xlsx")
IMPORT_MARKER = "/* === 2026-10-06 quote workbook import === */"


BRANDS = [
    {
        "name": "Wakodo",
        "kana": "和光堂",
        "letter": "W",
        "hue": "#D59A62",
        "blurb": "Asahi Group Foods · baby food and family care",
        "items": 140,
        "kind": "babykids",
        "site": "https://www.wakodo.co.jp/",
        "name_ja": "和光堂",
        "name_zh": "Wakodo 和光堂",
    },
    {
        "name": "Akai Bohshi",
        "kana": "赤い帽子",
        "letter": "A",
        "hue": "#A94335",
        "blurb": "Tivoli · Japanese cookies and gift assortments",
        "items": 21,
        "kind": "sweets",
        "site": "https://www.akaibohshi.com/jp/",
        "name_ja": "赤い帽子",
        "name_zh": "Akai Bohshi 红帽子",
    },
    {
        "name": "Moegino",
        "kana": "もえぎ野",
        "letter": "M",
        "hue": "#7B9360",
        "blurb": "Tivoli · Japanese baked confectionery assortments",
        "items": 5,
        "kind": "sweets",
        "site": "https://www.akaibohshi.com/jp/",
        "name_ja": "もえぎ野",
        "name_zh": "Moegino 萌木野",
    },
    {
        "name": "RUYSDAEL",
        "kana": "ロイスダール",
        "letter": "R",
        "hue": "#6F5846",
        "blurb": "Japanese confectionery and baked gifts",
        "items": 1,
        "kind": "sweets",
        "name_ja": "ロイスダール",
        "name_zh": "RUYSDAEL",
    },
]


def jan_text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value or "").strip()


def normalize_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("　", " ")).strip()


def parse_case_pack(value: object) -> int:
    if isinstance(value, (int, float)):
        return max(1, int(value))
    text = unicodedata.normalize("NFKC", str(value or ""))
    factors = [int(n) for n in re.findall(r"\d+", text)]
    if "x" in text.lower() or "×" in text:
        result = 1
        for factor in factors:
            result *= factor
        return max(1, result)
    return factors[0] if factors else 1


def infer_unit(name: str) -> str:
    normalized = unicodedata.normalize("NFKC", name)
    matches = re.findall(
        r"\d+(?:\.\d+)?\s*(?:kg|g|ml|l|枚|本|包|個|袋|食|錠|粒|回分|P)(?:\s*[x×]\s*\d+(?:P|本|個|袋)?)?",
        normalized,
        flags=re.IGNORECASE,
    )
    return matches[-1].replace(" ", "") if matches else "1 unit"


def classify_wakodo(name: str) -> tuple[str, str, str]:
    if any(word in name for word in ("珈琲", "ミルクティー")):
        return "food", "coffee", "Japanese instant beverage"
    if "シッカロール" in name or "服用ゼリー" in name:
        return "babykids", "babyhygiene", "Japanese baby and family care product"
    if any(word in name for word in ("おやつ", "ジュレ", "果実", "むぎ茶", "麦茶", "ほうじ茶", "アクアライト", "カルピス", "スティック", "シリアル")):
        return "babykids", "kidsdrink", "Japanese baby snack or drink"
    return "babykids", "babyfood", "Japanese baby food"


def product_copy(name: str, brand: str, jan: str, unit: str, moq: int, product_type: str) -> tuple[str, str]:
    description = (
        f"{name} by {brand}, JAN {jan}. {product_type} for wholesale sourcing from Japan; "
        f"catalog unit {unit}, case pack {moq}. Request an export quotation from JAPANITEM."
    )
    normalized_type = product_type.removeprefix("Japanese ").lower()
    lead = (
        f"{name} is a Japanese-market {normalized_type} listed for wholesale sourcing. "
        f"Use JAN / GTIN {jan}, unit {unit} and case pack {moq} to identify the exact SKU when requesting a quote."
    )
    return description, lead


def quote_rows() -> dict[str, dict]:
    sheet = openpyxl.load_workbook(QUOTE_FILE, data_only=True).active
    products: dict[str, dict] = {}
    for jan_value, raw_name, case_pack, price, _future_price in sheet.iter_rows(min_row=2, values_only=True):
        jan = jan_text(jan_value)
        name = normalize_name(raw_name)
        if not jan or not name or str(price or "").strip() == "販売終了":
            continue
        if jan.startswith("497518"):
            brand, category, sub, product_type = "Akai Bohshi", "sweets", "biscuit", "Japanese cookie"
        elif "RUYSDAEL" in name.upper():
            brand, category, sub, product_type = "RUYSDAEL", "sweets", "yogashi", "Japanese baked confectionery"
        else:
            brand = "Wakodo"
            category, sub, product_type = classify_wakodo(name)
        unit = infer_unit(name)
        moq = parse_case_pack(case_pack)
        seo_description, seo_lead = product_copy(name, brand, jan, unit, moq, product_type)
        products[jan] = {
            "id": jan,
            "jan": jan,
            "name": name,
            "brand": brand,
            "category": category,
            "sub": sub,
            "price": int(price) if isinstance(price, (int, float)) else 0,
            "moq": moq,
            "unit": unit,
            "tag": "new",
            "glyph": "子" if category == "babykids" else "甘",
            "hue": next(b["hue"] for b in BRANDS if b["name"] == brand),
            "seo_description": seo_description,
            "seo_lead": seo_lead,
        }
    return products


def akaibohshi_rows() -> tuple[dict[str, dict], dict[int, bytes]]:
    book = openpyxl.load_workbook(AKAIBOHSHI_FILE, data_only=True)
    sheet = book.active
    images = {image.anchor._from.row + 1: image._data() for image in sheet._images}
    products: dict[str, dict] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
        jan = jan_text(row[0])
        if not jan:
            continue
        raw_brand, raw_unit, raw_name, raw_pack = row[1], row[2], row[3], row[4]
        brand = "Akai Bohshi" if normalize_name(raw_brand) == "赤い帽子" else "Moegino"
        name = normalize_name(raw_name)
        unit = normalize_name(raw_unit) or infer_unit(name)
        moq = parse_case_pack(raw_pack)
        price = row[9] if isinstance(row[9], (int, float)) else 0
        product_type = "Japanese cookie or baked confectionery"
        seo_description, seo_lead = product_copy(name, brand, jan, unit, moq, product_type)
        products[jan] = {
            "id": jan,
            "jan": jan,
            "name": name,
            "brand": brand,
            "category": "sweets",
            "sub": "biscuit",
            "price": int(price),
            "moq": moq,
            "unit": unit,
            "tag": "new",
            "glyph": "甘",
            "hue": next(b["hue"] for b in BRANDS if b["name"] == brand),
            "shelf_life": normalize_name(row[8]),
            "seo_description": seo_description,
            "seo_lead": seo_lead,
            "_source_row": row_number,
        }
    return products, images


def export_images(products: dict[str, dict], images: dict[int, bytes]) -> int:
    PRODUCT_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    exported = 0
    for product in products.values():
        source = images.get(product.pop("_source_row"))
        if not source:
            continue
        image = Image.open(io.BytesIO(source)).convert("RGBA")
        background = Image.new("RGBA", image.size, "white")
        background.alpha_composite(image)
        rgb = background.convert("RGB")
        rgb.thumbnail((900, 900), Image.Resampling.LANCZOS)
        rgb.save(PRODUCT_IMAGE_DIR / f"{product['jan']}.jpg", "JPEG", quality=88, optimize=True, progressive=True)
        exported += 1
    return exported


def append_catalog(products: list[dict]) -> None:
    text = DATA_FILE.read_text(encoding="utf-8")
    if IMPORT_MARKER in text:
        raise SystemExit("The 2026-10-06 quote import is already present; no changes made.")

    brand_anchor = "\n];\n\nconst PRODUCTS = ["
    if brand_anchor not in text:
        raise RuntimeError("Could not locate BRANDS array boundary")
    brand_json = ",\n".join("  " + json.dumps(brand, ensure_ascii=False) for brand in BRANDS)
    text = text.replace(brand_anchor, f"\n{brand_json}\n];\n\nconst PRODUCTS = [", 1)

    product_anchor = "\n];\n\nwindow.CATEGORIES = CATEGORIES;"
    if product_anchor not in text:
        raise RuntimeError("Could not locate PRODUCTS array boundary")
    clean_products = [{key: value for key, value in p.items() if not key.startswith("_")} for p in products]
    product_json = ",\n".join("  " + json.dumps(product, ensure_ascii=False) for product in clean_products)
    text = text.replace(product_anchor, f",\n  {IMPORT_MARKER}\n{product_json}\n];\n\nwindow.CATEGORIES = CATEGORIES;", 1)
    DATA_FILE.write_text(text, encoding="utf-8")


def main() -> None:
    quote = quote_rows()
    akaibohshi, images = akaibohshi_rows()
    quote.update(akaibohshi)
    if len(quote) != 167:
        raise RuntimeError(f"Expected 167 unique active JANs, found {len(quote)}")
    exported = export_images(akaibohshi, images)
    append_catalog(list(quote.values()))
    print(json.dumps({"products_added": len(quote), "images_exported": exported}, ensure_ascii=False))


if __name__ == "__main__":
    main()
